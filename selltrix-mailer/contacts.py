"""Read only explicitly approved, unflagged Excel recipients."""
from collections.abc import Generator, Iterator
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path

from email_validator import EmailNotValidError, validate_email
from openpyxl import load_workbook


@dataclass(frozen=True, slots=True)
class InputError(Exception):
    detail: str

    def __str__(self) -> str:
        return self.detail


@dataclass(frozen=True, slots=True)
class RecipientBatch:
    path: Path
    allowed: int
    skipped: int

    def __iter__(self) -> Iterator[str]:
        return _iter_recipients(self.path)

    def __len__(self) -> int:
        return self.allowed


def address(value: str) -> str:
    try:
        return validate_email(
            value.strip(), check_deliverability=False, allow_smtputf8=False,
        ).normalized.lower()
    except EmailNotValidError as exc:
        raise InputError(f"Некорректный email: {value}") from exc


def _iter_recipients(path: Path) -> Generator[str, None, int]:
    """Require an explicit approval column; do not infer consent from names."""
    with closing(load_workbook(path, read_only=True, data_only=False)) as book:
        sheet_name = "К отправке"
        if sheet_name not in book.sheetnames:
            sheet_name = "Очищенный список"
        if sheet_name not in book.sheetnames:
            raise InputError("Нужен лист «К отправке» или «Очищенный список».")
        rows = book[sheet_name].iter_rows(values_only=True)
        header = [str(v or "").strip().lower() for v in next(rows, ())]
        email_key = "email" if "email" in header else "почта"
        if email_key not in header or "approved" not in header:
            raise InputError("Добавьте столбец approved: ДА только для согласованных получателей.")
        email_i, approved_i = header.index(email_key), header.index("approved")
        flag_i = header.index("флаги") if "флаги" in header else None
        blocked: set[str] = set()
        for name in ("Исправления доменов", "Требует проверки", "Полный аудит изменений"):
            if name not in book.sheetnames:
                continue
            for row in book[name].iter_rows(min_row=2, values_only=True):
                if len(row) >= 7 and (name != "Полный аудит изменений" or "R3" in str(row[6])):
                    blocked.add(str(row[4] or "").strip().lower())
        seen: set[str] = set()
        skipped = 0
        for row in rows:
            raw = str(row[email_i] or "").strip()
            approved = str(row[approved_i] or "").strip().upper() == "ДА"
            flagged = flag_i is not None and bool(row[flag_i])
            try:
                email = address(raw)
            except InputError:
                skipped += 1
                continue
            if not approved or flagged or email in blocked or email in seen:
                skipped += 1
                continue
            seen.add(email)
            yield email
        return skipped


def recipients(path: Path) -> RecipientBatch:
    source = _iter_recipients(path)
    allowed = 0
    while True:
        try:
            next(source)
            allowed += 1
        except StopIteration as result:
            return RecipientBatch(path, allowed, int(result.value))
